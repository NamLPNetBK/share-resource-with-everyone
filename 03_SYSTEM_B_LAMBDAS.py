"""
SYSTEM B: USER UPLOAD EVALUATION
- User uploads new project estimate
- Compare with standards calculated from dataset
- AI evaluation (using Bedrock)
- Assessment and recommendations
"""

import boto3
import json
import openpyxl
from io import BytesIO
from datetime import datetime
from decimal import Decimal

s3_client = boto3.client('s3')
dynamodb = boto3.resource('dynamodb')
eventbridge_client = boto3.client('events')
bedrock_client = boto3.client('bedrock-runtime', region_name='us-east-1')


# ====== LAMBDA 1: ParseUserUpload ======

def parse_user_project(file_content):
    """
    Parse user-uploaded project file
    Same format as dataset but for a single project
    
    Returns:
    {
      'summary': {...},
      'items': [...]
    }
    """
    wb = openpyxl.load_workbook(BytesIO(file_content))
    ws = wb.active
    
    # Parse summary (rows 1-5)
    summary = {}
    for row_idx in range(1, 6):
        cell_a = ws[f'A{row_idx}'].value
        cell_b = ws[f'B{row_idx}'].value
        
        if cell_a == '案件番號':
            summary['projectNumber'] = str(cell_b)
        elif cell_a == '總工數':
            summary['totalWorkload'] = float(cell_b)
        elif cell_a == '畫面總數':
            summary['totalPage'] = int(cell_b)
        elif cell_a == 'API總數':
            summary['totalApi'] = int(cell_b)
        elif cell_a == '案件對應日':
            summary['projectDay'] = str(cell_b)
    
    # Parse items (rows 7+)
    items = []
    for row_idx in range(7, ws.max_row + 1):
        item_id = ws[f'A{row_idx}'].value
        name = ws[f'B{row_idx}'].value
        class_type = ws[f'C{row_idx}'].value
        difficulty = ws[f'D{row_idx}'].value
        workload = ws[f'E{row_idx}'].value
        
        if not item_id:
            break
        
        items.append({
            'itemId': str(item_id),
            'name': str(name),
            'class': class_type,
            'difficulty': difficulty,
            'workload': float(workload)
        })
    
    return {'summary': summary, 'items': items}


def lambda_parse_user_upload(event, context):
    """
    Lambda 1: Parse user-uploaded file
    Trigger: API Gateway POST /api/analyze
    
    Event:
    {
      "execution_id": "exec-user-001",
      "bucket": "user-uploads",
      "file_key": "PRJ-USER-001.xlsx"
    }
    """
    try:
        execution_id = event['execution_id']
        bucket = event['bucket']
        file_key = event['file_key']
        
        print(f"[ParseUser] Processing: {execution_id}")
        
        workflow_table = dynamodb.Table('WorkflowExecutions')
        
        # Update status
        workflow_table.update_item(
            Key={'execution_id': execution_id},
            UpdateExpression='SET #status = :status, #step = :step, updated_at = :now',
            ExpressionAttributeNames={
                '#status': 'status',
                '#step': 'current_step'
            },
            ExpressionAttributeValues={
                ':status': 'PARSING',
                ':step': 'Parsing user upload...',
                ':now': datetime.now().isoformat()
            }
        )
        
        # Download and parse
        obj = s3_client.get_object(Bucket=bucket, Key=file_key)
        file_content = obj['Body'].read()
        
        parsed = parse_user_project(file_content)
        
        print(f"[ParseUser] Parsed {len(parsed['items'])} items")
        
        # Publish event
        eventbridge_client.put_events(
            Entries=[{
                'Source': 'workload.user',
                'DetailType': 'UserUploadParsed',
                'Detail': json.dumps({
                    'execution_id': execution_id,
                    'parsed_data': parsed
                }, default=str)
            }]
        )
        
        return {'statusCode': 200, 'execution_id': execution_id}
        
    except Exception as e:
        print(f"[ParseUser] Error: {str(e)}")
        return {'statusCode': 500, 'error': str(e)}


# ====== LAMBDA 2: AnalyzeVariance ======

def lambda_analyze_variance(event, context):
    """
    Lambda 2: Analyze variance against standards
    Trigger: EventBridge (UserUploadParsed)
    
    - Get standards from 工數基準表
    - Calculate variance for each item
    - Identify high variance items
    """
    try:
        execution_id = event['detail']['execution_id']
        parsed_data = json.loads(event['detail']['parsed_data'])
        
        print(f"[AnalyzeVariance] Processing: {execution_id}")
        
        workflow_table = dynamodb.Table('WorkflowExecutions')
        standards_table = dynamodb.Table('工數基準表')
        
        # Update status
        workflow_table.update_item(
            Key={'execution_id': execution_id},
            UpdateExpression='SET #status = :status, #step = :step',
            ExpressionAttributeNames={
                '#status': 'status',
                '#step': 'current_step'
            },
            ExpressionAttributeValues={
                ':status': 'ANALYZING',
                ':step': 'Analyzing variance...'
            }
        )
        
        # Analyze each item
        items = parsed_data['items']
        analyzed_items = []
        
        ok_count = 0
        warning_count = 0
        alert_count = 0
        total_variance = 0
        
        for item in items:
            # Get standard
            try:
                response = standards_table.get_item(
                    Key={
                        'class': item['class'],
                        'difficulty': item['difficulty']
                    }
                )
                
                standard = float(response['Item']['averageWorkload']) if 'Item' in response else 0
            except:
                standard = 0
            
            if standard == 0:
                variance = 0
                variance_percent = 0
                status = 'UNKNOWN'
            else:
                variance = item['workload'] - standard
                variance_percent = (variance / standard * 100)
                
                if abs(variance_percent) <= 20:
                    status = 'OK'
                    ok_count += 1
                elif abs(variance_percent) <= 50:
                    status = 'WARNING'
                    warning_count += 1
                else:
                    status = 'ALERT'
                    alert_count += 1
            
            total_variance += abs(variance_percent)
            
            analyzed_items.append({
                'itemId': item['itemId'],
                'name': item['name'],
                'class': item['class'],
                'difficulty': item['difficulty'],
                'actual': item['workload'],
                'standard': standard,
                'variance': round(variance, 2),
                'variancePercent': round(variance_percent, 2),
                'status': status
            })
        
        avg_variance = total_variance / len(items) if items else 0
        
        # Save analysis result
        analysis_result = {
            'itemVariances': analyzed_items,
            'summary': {
                'ok_count': ok_count,
                'warning_count': warning_count,
                'alert_count': alert_count,
                'average_variance_percent': round(avg_variance, 2)
            }
        }
        
        # Update workflow
        workflow_table.update_item(
            Key={'execution_id': execution_id},
            UpdateExpression='SET analysis_result = :result, #step = :step',
            ExpressionAttributeNames={'#step': 'current_step'},
            ExpressionAttributeValues={
                ':result': analysis_result,
                ':step': 'Variance analysis complete, preparing AI evaluation...'
            }
        )
        
        print(f"[AnalyzeVariance] OK={ok_count}, WARNING={warning_count}, ALERT={alert_count}")
        
        # Publish event
        eventbridge_client.put_events(
            Entries=[{
                'Source': 'workload.user',
                'DetailType': 'VarianceAnalyzed',
                'Detail': json.dumps({
                    'execution_id': execution_id,
                    'analysis_result': analysis_result,
                    'parsed_data': parsed_data
                }, default=str)
            }]
        )
        
        return {'statusCode': 200}
        
    except Exception as e:
        print(f"[AnalyzeVariance] Error: {str(e)}")
        import traceback
        traceback.print_exc()
        return {'statusCode': 500, 'error': str(e)}


# ====== LAMBDA 3: AIEvaluation ======

def build_ai_prompt(parsed_data, analysis_result):
    """
    Build prompt for Bedrock Claude
    Include: project data, variances, high-variance items
    """
    
    project = parsed_data['summary']
    items = analysis_result['itemVariances']
    summary = analysis_result['summary']
    
    # Filter high variance items
    high_variance = [i for i in items if i['status'] in ['WARNING', 'ALERT']]
    high_variance.sort(key=lambda x: abs(x['variancePercent']), reverse=True)
    
    prompt = f"""
あなたはシステム開発プロジェクトの工数見積もり専門家です。
以下のプロジェクトの工数見積もり評価を実施してください。

【プロジェクト情報】
案件番號: {project.get('projectNumber', 'N/A')}
總工數: {project.get('totalWorkload', 0)} 人日
畫面總數: {project.get('totalPage', 0)}
API總數: {project.get('totalApi', 0)}
案件對應日: {project.get('projectDay', 'N/A')}

【統計情報】
工數項目數: {len(items)}
OK項目: {summary['ok_count']}
注意項目: {summary['warning_count']}
警告項目: {summary['alert_count']}
平均乖離率: {summary['average_variance_percent']:.2f}%

【乖離の大きい項目トップ5】
"""
    
    for item in high_variance[:5]:
        prompt += f"\n- {item['name']} ({item['class']}/{item['difficulty']})"
        prompt += f"\n  見積: {item['actual']:.1f} 人日, 標準: {item['standard']:.1f} 人日"
        prompt += f"\n  乖離: {item['variance']:+.1f} ({item['variancePercent']:+.1f}%)"

    prompt += f"""

【評価項目】
1. 全体的な工数見積もりの妥當性
2. 乖離が大きい理由の推測
3. 最適化の余地の有無
4. リスク要因
5. 改善提案

日本語で詳細に評価してください（3-4段落で簡潔に）。
"""
    
    return prompt


def lambda_ai_evaluation(event, context):
    """
    Lambda 3: AI evaluation using Bedrock
    Trigger: EventBridge (VarianceAnalyzed)
    
    - Build prompt with project data and variances
    - Call Bedrock Claude
    - Get evaluation in Japanese
    - Determine assessment level
    """
    try:
        execution_id = event['detail']['execution_id']
        analysis_result = json.loads(event['detail']['analysis_result'])
        parsed_data = json.loads(event['detail']['parsed_data'])
        
        print(f"[AIEvaluation] Processing: {execution_id}")
        
        workflow_table = dynamodb.Table('WorkflowExecutions')
        
        # Update status
        workflow_table.update_item(
            Key={'execution_id': execution_id},
            UpdateExpression='SET #status = :status, #step = :step',
            ExpressionAttributeNames={
                '#status': 'status',
                '#step': 'current_step'
            },
            ExpressionAttributeValues={
                ':status': 'EVALUATING_AI',
                ':step': 'Calling AI for evaluation...'
            }
        )
        
        # Build prompt
        prompt = build_ai_prompt(parsed_data, analysis_result)
        
        print("[AIEvaluation] Calling Bedrock Claude...")
        
        # Call Bedrock
        response = bedrock_client.invoke_model(
            modelId='anthropic.claude-3-5-sonnet-20241022',
            contentType='application/json',
            accept='application/json',
            body=json.dumps({
                'anthropic_version': 'bedrock-2023-06-01',
                'max_tokens': 1024,
                'messages': [
                    {
                        'role': 'user',
                        'content': prompt
                    }
                ]
            })
        )
        
        # Parse response
        result_body = json.loads(response['body'].read())
        ai_evaluation = result_body['content'][0]['text']
        
        print(f"[AIEvaluation] Received evaluation ({len(ai_evaluation)} chars)")
        
        # Determine assessment
        summary = analysis_result['summary']
        avg_variance = summary['average_variance_percent']
        alert_count = summary['alert_count']
        
        if alert_count > 0 or avg_variance > 30:
            assessment = 'NEEDS_REVIEW'
            assessment_label = '要檢討'
        elif summary['warning_count'] > 2 or avg_variance > 20:
            assessment = 'CAUTION'
            assessment_label = '注意'
        else:
            assessment = 'ACCEPTABLE'
            assessment_label = '妥當'
        
        # Save to AnalysisResults
        results_table = dynamodb.Table('AnalysisResults')
        
        final_result = {
            'execution_id': execution_id,
            'projectNumber': parsed_data['summary']['projectNumber'],
            'analysisDate': datetime.now().isoformat(),
            'projectData': parsed_data['summary'],
            'variance_analysis': analysis_result,
            'ai_evaluation': {
                'overallAssessment': assessment,
                'assessmentLabel': assessment_label,
                'evaluation': ai_evaluation
            }
        }
        
        results_table.put_item(Item=final_result)
        
        # Update workflow to COMPLETED
        workflow_table.update_item(
            Key={'execution_id': execution_id},
            UpdateExpression='SET #status = :status, #step = :step, ai_evaluation = :eval',
            ExpressionAttributeNames={
                '#status': 'status',
                '#step': 'current_step'
            },
            ExpressionAttributeValues={
                ':status': 'COMPLETED',
                ':step': f'Evaluation complete - Assessment: {assessment_label}',
                ':eval': {
                    'assessment': assessment,
                    'label': assessment_label
                }
            }
        )
        
        print(f"[AIEvaluation] Assessment: {assessment_label}")
        
        # Publish event
        eventbridge_client.put_events(
            Entries=[{
                'Source': 'workload.user',
                'DetailType': 'EvaluationComplete',
                'Detail': json.dumps({
                    'execution_id': execution_id,
                    'assessment': assessment,
                    'assessment_label': assessment_label
                })
            }]
        )
        
        return {'statusCode': 200, 'assessment': assessment_label}
        
    except Exception as e:
        print(f"[AIEvaluation] Error: {str(e)}")
        import traceback
        traceback.print_exc()
        return {'statusCode': 500, 'error': str(e)}


# ====== LAMBDA 4: GetWorkflowStatus ======

def lambda_get_status(event, context):
    """
    Lambda 4: Get workflow execution status
    Trigger: API Gateway GET /api/status?execution_id=xxx
    
    Returns current status and results when complete
    """
    try:
        execution_id = event['queryStringParameters']['execution_id']
        
        workflow_table = dynamodb.Table('WorkflowExecutions')
        results_table = dynamodb.Table('AnalysisResults')
        
        # Get workflow status
        response = workflow_table.get_item(Key={'execution_id': execution_id})
        
        if 'Item' not in response:
            return {
                'statusCode': 404,
                'body': json.dumps({'error': 'Execution not found'})
            }
        
        workflow = response['Item']
        status = workflow['status']
        
        result = {
            'execution_id': execution_id,
            'status': status,
            'current_step': workflow.get('current_step', ''),
            'updated_at': workflow.get('updated_at', '')
        }
        
        # If completed, get full results
        if status == 'COMPLETED':
            try:
                result_response = results_table.get_item(
                    Key={'execution_id': execution_id}
                )
                
                if 'Item' in result_response:
                    analysis = result_response['Item']
                    result['analysis'] = {
                        'projectNumber': analysis['projectNumber'],
                        'projectData': analysis['projectData'],
                        'variance_analysis': analysis['variance_analysis'],
                        'ai_evaluation': analysis['ai_evaluation']
                    }
            except:
                pass
        
        return {
            'statusCode': 200,
            'body': json.dumps(result, default=str)
        }
        
    except Exception as e:
        print(f"[GetStatus] Error: {str(e)}")
        return {'statusCode': 500, 'body': json.dumps({'error': str(e)})}
